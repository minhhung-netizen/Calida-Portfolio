"""Đọc/ghi sheet Excel theo SCHEMA, upsert theo khóa (không ghi đè lịch sử)."""
from pathlib import Path
from uuid import uuid4
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from schema import SCHEMA
from config import INPUT_DIR


NUMBER_FORMATS = {
    "open": "#,##0.00;[Red](#,##0.00);-", "high": "#,##0.00;[Red](#,##0.00);-",
    "low": "#,##0.00;[Red](#,##0.00);-", "close": "#,##0.00;[Red](#,##0.00);-",
    "support_lo": "#,##0.00;[Red](#,##0.00);-", "support_hi": "#,##0.00;[Red](#,##0.00);-",
    "resist_lo": "#,##0.00;[Red](#,##0.00);-", "resist_hi": "#,##0.00;[Red](#,##0.00);-",
    "expected_lo": "#,##0.00;[Red](#,##0.00);-", "expected_hi": "#,##0.00;[Red](#,##0.00);-",
    "buy_lo": "#,##0.00;[Red](#,##0.00);-", "buy_hi": "#,##0.00;[Red](#,##0.00);-",
    "target": "#,##0.00;[Red](#,##0.00);-", "stop": "#,##0.00;[Red](#,##0.00);-",
    "cost": "#,##0.00;[Red](#,##0.00);-", "nav_bn": "#,##0.00;[Red](#,##0.00);-",
    "vn_target": "#,##0;[Red](#,##0);-", "volume": "#,##0;[Red](#,##0);-",
    "net_value": "#,##0.00;[Red](#,##0.00);-",
    "chg_pct": "0.0;[Red](0.0);-", "weight_pct": "0.0;[Red](0.0);-",
    "ytd_pct": "0.0;[Red](0.0);-", "stock_pct": "0.0;[Red](0.0);-",
    "cash_pct": "0.0;[Red](0.0);-", "other_pct": "0.0;[Red](0.0);-",
}


def spec(file: str, sheet: str) -> dict:
    return SCHEMA[file][sheet]


def read_sheet(file: str, sheet: str) -> pd.DataFrame:
    s = spec(file, sheet)
    path = INPUT_DIR / file
    if not path.exists():
        return pd.DataFrame(columns=s["cols"])
    try:
        df = pd.read_excel(path, sheet_name=sheet)
    except ValueError:  # sheet chưa có
        return pd.DataFrame(columns=s["cols"])
    df.columns = [str(c).strip() for c in df.columns]
    return df


def write_sheets(file: str, frames: dict):
    """Ghi nhiều sheet vào 1 file, giữ nguyên các sheet khác đã có."""
    path = INPUT_DIR / file
    path.parent.mkdir(parents=True, exist_ok=True)
    existing = {}
    if path.exists():
        existing = pd.read_excel(path, sheet_name=None)
    for sh in SCHEMA[file]:
        if sh not in frames and sh not in existing:
            existing[sh] = pd.DataFrame(columns=SCHEMA[file][sh]["cols"])
    existing.update(frames)
    order = list(SCHEMA[file]) + [k for k in existing if k not in SCHEMA[file]]
    tmp = path.with_name(f".{path.stem}-{uuid4().hex}.tmp{path.suffix}")
    try:
        with pd.ExcelWriter(tmp, engine="openpyxl") as w:
            for sh in order:
                df = existing[sh]
                cols = SCHEMA[file].get(sh, {}).get("cols")
                if cols:
                    df = df.reindex(columns=cols)
                df.to_excel(w, sheet_name=sh, index=False)
        _style(tmp)
        tmp.replace(path)
    finally:
        if tmp.exists():
            tmp.unlink()


def upsert(file: str, sheet: str, new: pd.DataFrame):
    s = spec(file, sheet)
    old = read_sheet(file, sheet)
    new = new.reindex(columns=s["cols"])
    df = pd.concat([old.reindex(columns=s["cols"]), new], ignore_index=True)
    for c in s.get("dates", []):
        from build_db import to_date
        df[c] = to_date(df[c])
    df = df.drop_duplicates(subset=s["key"], keep="last")
    sort = [c for c in s["key"] if c in df.columns]
    df = df.sort_values(sort).reset_index(drop=True)
    write_sheets(file, {sheet: df})
    return len(df)


def _style(path: Path):
    wb = load_workbook(path)
    head = PatternFill("solid", fgColor="1F5EEA")
    for ws in wb.worksheets:
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = head
        ws.freeze_panes = "A2"
        for column in ws[1]:
            if column.value in NUMBER_FORMATS:
                for cell in ws.iter_cols(min_col=column.column, max_col=column.column, min_row=2):
                    for value in cell:
                        value.number_format = NUMBER_FORMATS[column.value]
        for i, col in enumerate(ws.columns, 1):
            width = max((len(str(c.value)) if c.value is not None else 0) for c in list(col)[:200])
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)
    wb.save(path)
